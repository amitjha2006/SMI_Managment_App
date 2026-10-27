"""Import reports persist row-level reasons beyond the on-screen result.

Run: python -m pytest backend/tests/test_import_reports.py -v
"""

import asyncio
import io
import os
import sys
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

import backend.main as m  # noqa: E402


def _req():
    return SimpleNamespace(headers={}, client=None)


def _supplier_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Supplier Name", "Contact Number"])
    ws.append(["Good Supplier", "9000000091"])
    ws.append(["Good Supplier", "9000000091"])  # skipped: already exists
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _run_import(db, monkeypatch, content, filename="s.xlsx"):
    def fake_backup(*_args, **_kwargs):
        return SimpleNamespace(id=99, filename="pre-import_test.db")

    monkeypatch.setattr(m, "_do_backup", fake_backup)

    async def fake_read():
        return content

    return asyncio.run(m.import_data(
        "suppliers", _req(),
        SimpleNamespace(filename=filename, read=fake_read), None, db))


def test_report_keeps_skipped_reasons(db, monkeypatch):
    out = _run_import(db, monkeypatch, _supplier_xlsx())
    assert out["imported"] == 1, out
    assert out["skipped"] == 1, out
    rows = m.list_import_reports(db=db)
    assert len(rows) == 1, rows
    rep = rows[0]
    assert rep["filename"] == "s.xlsx", rep
    assert rep["imported"] == 1 and rep["skipped"] == 1, rep
    reasons = " ".join(r["reason"] for r in rep["skipped_rows"])
    assert "already exists" in reasons, rep["skipped_rows"]


def test_reports_pruned_to_twenty(db, monkeypatch):
    for i in range(22):
        _run_import(db, monkeypatch, _supplier_xlsx(), filename=f"s{i}.xlsx")
    rows = m.list_import_reports(limit=50, db=db)
    assert len(rows) == 20, len(rows)
    assert rows[0]["filename"] == "s21.xlsx", rows[0]
    assert rows[-1]["filename"] == "s2.xlsx", rows[-1]


def test_report_survives_datetime_and_decimal_row_data(db, monkeypatch):
    # Skipped-row data carries raw Excel values (datetimes, Decimals):
    # the report store must serialize them like _audit does (default=str).
    from datetime import datetime

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Challan Number", "Challan Date", "Customer Name",
               "Product Name", "Packaging Size", "Quantity", "Rate per Kg"])
    ws.append(["CH-DT-1", datetime(2025, 4, 2), "Nobody Here",
               "Some Masala", "250g", 10, 200])
    buf = io.BytesIO()
    wb.save(buf)

    def fake_backup(*_args, **_kwargs):
        return SimpleNamespace(id=95, filename="pre-import_test.db")

    monkeypatch.setattr(m, "_do_backup", fake_backup)

    from backend.app import models as _models  # noqa: E402
    db.add(_models.Customer(customer_id="CUST-DT1", name="Somebody Else",
                            contact_number="9000000098"))
    db.commit()

    async def fake_read():
        return buf.getvalue()

    out = asyncio.run(m.import_data(
        "sales_challans", _req(),
        SimpleNamespace(filename="c.xlsx", read=fake_read), None, db))
    assert out["status"] == "success", out
    assert out["skipped"] == 1, out
    rows = m.list_import_reports(db=db)
    assert len(rows) == 1, rows
    assert "Customer not found" in " ".join(
        r["reason"] for r in rows[0]["skipped_rows"])


def test_reports_page_lists_reasons():
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    with open(os.path.join(root, "frontend", "pages", "import.html"), encoding="utf-8") as fh:
        html = fh.read()
    assert 'id="report-sel"' in html
    assert "/import/reports" in html
    # Both upload paths (preview + import) read text first so a plain-text
    # 500 surfaces its message instead of a JSON parse error.
    assert html.count("await res.text()") == 2


def test_result_persistence_failure_stays_json(db, monkeypatch):
    import asyncio
    from types import SimpleNamespace as _NS

    def fake_backup(*_args, **_kwargs):
        return _NS(id=96, filename="pre-import_test.db")

    def boom_save(*_args, **_kwargs):
        raise RuntimeError("report store down")

    monkeypatch.setattr(m, "_do_backup", fake_backup)
    monkeypatch.setattr(m, "_save_import_report", boom_save)

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Supplier Name", "Contact Number"])
    ws.append(["Report Down Supplier", "9000000099"])
    buf = io.BytesIO()
    wb.save(buf)

    async def fake_read():
        return buf.getvalue()

    with pytest.raises(HTTPException) as exc:
        asyncio.run(m.import_data(
            "suppliers", _req(),
            SimpleNamespace(filename="s.xlsx", read=fake_read), None, db))
    assert exc.value.status_code == 500
    assert "recording the result failed" in exc.value.detail


def test_import_parses_plain_text_errors():
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    with open(os.path.join(root, "frontend", "pages", "import.html"), encoding="utf-8") as fh:
        html = fh.read()
    # Both upload paths (preview + import) read text first so a plain-text
    # 500 surfaces its message instead of a JSON parse error.
    assert html.count("await res.text()") == 2


def test_report_csv_download_wired():
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    with open(os.path.join(root, "frontend", "pages", "import.html"), encoding="utf-8") as fh:
        html = fh.read()
    assert "downloadReportCsv" in html
    assert "Download CSV" in html
    with open(os.path.join(root, "frontend", "js", "components.js"), encoding="utf-8") as fh:
        js = fh.read()
    assert "downloadBlob" in js


def test_cards_fail_visible_with_retry():
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    with open(os.path.join(root, "frontend", "pages", "import.html"), encoding="utf-8") as fh:
        html = fh.read()
    # Load failures must name themselves (never a silent blank) and retry once.
    assert "Sign in again or press Refresh" in html
    assert "loadReports(true)" in html
    assert "loadUndoStatus(true)" in html


def test_reports_load_on_page_entry():
    # Regression: the Initialize block once forgot loadReports(), so the
    # card only filled right after a same-session import/undo and showed
    # blank on every navigation or refresh.
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    with open(os.path.join(root, "frontend", "pages", "import.html"), encoding="utf-8") as fh:
        html = fh.read()
    init = html.split("// Initialize", 1)[1]
    assert "loadReports();" in init, "page entry never loads past reports"
