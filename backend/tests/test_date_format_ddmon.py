"""One date format everywhere: DD-Mon-YYYY ("05-Mar-2026").

A day-first numeric form (05/06/2026) can still be misread across
locales; a short month name cannot. Every human-facing date renders
through this shape - table cells, print views, PDFs, messages, Excel -
while machine formats (API payloads, filenames, logs, DB storage,
period codes, chart keys) stay ISO.

Run: python -m pytest backend/tests/test_date_format_ddmon.py -v
"""

import os
import re


def _read(rel):
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    with open(os.path.join(root, *rel.split("/")), encoding="utf-8") as fh:
        return fh.read()


def test_shared_helpers_pin_dd_mon():
    js = _read("frontend/js/components.js")
    i = js.find("function fmtDay")
    assert i >= 0, "no shared UI.fmtDay helper"
    body = js[i:i + 600]
    assert '"en-IN"' in body and '"short"' in body, (
        "UI.fmtDay does not pin the DD-Mon render"
    )
    assert "fmtDay" in js.split("return {")[-1], "UI.fmtDay not exported"
    for helper in ("function dateWithTime", "function dateTimeText"):
        j = js.find(helper)
        assert j >= 0, f"{helper} missing"
        assert "fmtDay(" in js[j:j + 1200], (
            f"{helper} does not route through the shared DD-Mon render"
        )


def test_no_page_renders_bare_locale_dates():
    import glob as _glob

    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    offenders = []
    for path in sorted(_glob.glob(os.path.join(root, "frontend", "pages", "*.html"))):
        with open(path, encoding="utf-8") as fh:
            src = fh.read()
        # toLocaleDateString() with no locale, or with an undefined one,
        # renders MM/DD/YYYY on en-US machines and DD/MM/YYYY on en-IN
        # ones - never uniform. toLocaleString() bare has the same flaw.
        if (re.search(r"\.toLocaleDateString\(\s*\)", src)
                or re.search(r"\.toLocaleDateString\(undefined", src)
                or re.search(r"\.toLocaleString\(\)", src)):
            offenders.append(os.path.basename(path))
    assert not offenders, f"pages with locale-dependent date renders: {offenders}"


def test_invoice_pdf_shows_dd_mon():
    from backend.app.services import pdf_export as pe
    from backend.tests.test_invoice_pdf import _inv_dict, _pdf_text

    text = _pdf_text(pe.render_invoice_pdf(_inv_dict()))
    assert "28-Sep-2026" in text, "invoice PDF date is not DD-Mon-YYYY"


def test_import_reads_dd_mon():
    from backend.app.services.excel_import import parse_date

    got = parse_date("05-Mar-2026")
    assert got is not None and (got.year, got.month, got.day) == (2026, 3, 5)


def test_excel_date_cells_carry_dd_mmm_format():
    import io as _io
    from datetime import datetime

    from openpyxl import load_workbook

    from backend.app.services import excel_export as ee

    blob = ee._workbook("T", ["Date"], [{"Date": datetime(2026, 3, 5, 10, 0, 0)}])
    ws = load_workbook(filename=_io.BytesIO(blob)).active
    assert ws["A2"].number_format == "DD-MMM-YYYY", ws["A2"].number_format
