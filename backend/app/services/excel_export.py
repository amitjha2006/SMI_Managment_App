"""Excel export service for the Shresth Masala Industry Management System.

Exports database records to .xlsx files whose columns match the import
templates one-to-one, so an exported file can be re-imported as-is
(round-trip safe). Mirrors backend/app/services/excel_import.py.
"""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from backend.app.models import (
    Customer,
    CustomerPayment,
    Employee,
    Expense,
    PackagingSize,
    Product,
    PurchaseInvoice,
    RawMaterial,
    SalesChallan,
    Supplier,
)

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")


def _v(value: Any) -> Any:
    """Convert DB values to Excel-safe cell values."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, Decimal):
        return float(value)
    # Enum members (model and schema enums alike) -> their value
    if hasattr(value, "value") and not isinstance(value, (str, int, float)):
        try:
            return value.value
        except Exception:
            return str(value)
    return value


def _title(value: Any) -> Any:
    """full_time -> Full Time (round-trips through the fuzzy import parser)."""
    if value is None:
        return None
    text = str(getattr(value, "value", value))
    return text.replace("_", " ").title()


def _workbook(title: str, headers: list[str], rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r, row in enumerate(rows, 2):
        for col, header in enumerate(headers, 1):
            value = _v(row.get(header))
            cell = ws.cell(row=r, column=col, value=value)
            if isinstance(value, datetime):
                cell.number_format = "DD-MMM-YYYY"
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_customers(db: Session):
    headers = [
        "Customer ID",
        "Customer Name",
        "Company Name",
        "GST No",
        "PAN No",
        "Contact Number",
        "Contact Number 2",
        "Email",
        "Address",
        "Bank Account",
        "IFSC",
        "Bank Name",
        "Opening Balance",
        "Outstanding Balance",
        "Notes",
    ]
    rows = [
        {
            "Customer ID": c.customer_id,
            "Customer Name": c.name,
            "Company Name": c.company_name,
            "GST No": c.gst_no,
            "PAN No": c.pan_no,
            "Contact Number": c.contact_number,
            "Contact Number 2": c.contact_number_2,
            "Email": c.email,
            "Address": c.address,
            "Bank Account": c.account_no,
            "IFSC": c.ifsc_code,
            "Bank Name": c.bank_detail,
            "Opening Balance": c.opening_balance,
            "Outstanding Balance": c.outstanding_balance,
            "Notes": c.notes,
        }
        for c in db.query(Customer).order_by(Customer.name).all()
    ]
    return "Customers", headers, rows


def export_suppliers(db: Session):
    headers = [
        "Supplier ID",
        "Supplier Name",
        "Company Name",
        "GST No",
        "PAN No",
        "Contact Number",
        "Contact Number 2",
        "Email",
        "Address",
        "Bank Account",
        "IFSC",
        "Bank Name",
        "Opening Balance",
        "Outstanding Balance",
        "Notes",
    ]
    rows = [
        {
            "Supplier ID": s.supplier_id,
            "Supplier Name": s.name,
            "Company Name": s.company_name,
            "GST No": s.gst_no,
            "PAN No": s.pan_no,
            "Contact Number": s.contact_number,
            "Contact Number 2": s.contact_number_2,
            "Email": s.email,
            "Address": s.address,
            "Bank Account": s.account_no,
            "IFSC": s.ifsc_code,
            "Bank Name": s.bank_detail,
            "Opening Balance": s.opening_balance,
            "Outstanding Balance": s.outstanding_balance,
            "Notes": s.notes,
        }
        for s in db.query(Supplier).order_by(Supplier.name).all()
    ]
    return "Suppliers", headers, rows


def export_raw_materials(db: Session):
    headers = [
        "Raw Material ID",
        "Raw Material Name",
        "Category",
        "Unit",
        "Current Stock",
        "Opening Stock",
        "Threshold Stock",
        "Purchase Rate",
        "HSN Code",
    ]
    rows = [
        {
            "Raw Material ID": m.material_id,
            "Raw Material Name": m.name,
            "Category": m.category.name if m.category else None,
            "Unit": str(getattr(m.unit, "value", m.unit) or "").upper() or None,
            "Current Stock": m.current_stock,
            "Opening Stock": m.opening_stock,
            "Threshold Stock": m.threshold_stock,
            "Purchase Rate": m.purchase_rate,
            "HSN Code": m.hsn_code,
        }
        for m in db.query(RawMaterial).order_by(RawMaterial.name).all()
    ]
    return "Raw Materials", headers, rows


def export_products(db: Session):
    headers = [
        "Product ID",
        "Product Name",
        "Category",
        "Description",
        "Rate per Kg",
        "Packaging Size",
        "Weight (kg)",
        "HSN Code",
    ]
    rows = []
    for p in db.query(Product).order_by(Product.name).all():
        sizes = p.packaging_sizes or [None]
        for ps in sizes:
            rows.append(
                {
                    "Product ID": p.product_id,
                    "Product Name": p.name,
                    "Category": p.category,
                    "Description": p.description,
                    "Rate per Kg": p.base_rate_per_kg,
                    "Packaging Size": ps.size_name if ps else None,
                    "Weight (kg)": float(ps.weight_in_kg) if ps and ps.weight_in_kg else None,
                    "HSN Code": p.hsn_code,
                }
            )
    return "Products", headers, rows


def export_purchase_invoices(db: Session):
    headers = [
        "Invoice Number",
        "Invoice Date",
        "Supplier Name",
        "Raw Material Name",
        "Unit",
        "Quantity",
        "Rate",
        "GST %",
        "Paid Amount",
    ]
    rows = []
    for inv in db.query(PurchaseInvoice).order_by(PurchaseInvoice.invoice_date).all():
        items = inv.items or [None]
        for it in items:
            rows.append(
                {
                    "Invoice Number": inv.invoice_number,
                    "Invoice Date": inv.invoice_date,
                    "Supplier Name": inv.supplier.name if inv.supplier else None,
                    "Raw Material Name": it.raw_material.name if it and it.raw_material else None,
                    "Unit": it.unit if it else None,
                    "Quantity": it.quantity if it else None,
                    "Rate": it.rate if it else None,
                    "GST %": it.gst_percent if it else None,
                    "Paid Amount": inv.paid_amount,
                }
            )
    return "Purchase Invoices", headers, rows


def export_sales_challans(db: Session):
    headers = [
        "Challan Number",
        "Challan Date",
        "Customer Name",
        "Product Name",
        "Packaging Size",
        "Quantity",
        "Rate per Kg",
        "Amount",
        "Paid Amount",
        "Payment Mode",
    ]
    rows = []
    for ch in db.query(SalesChallan).order_by(SalesChallan.challan_date).all():
        pays = ch.payments or []
        mode = pays[0].payment_mode if pays else None
        if isinstance(mode, str):
            mode = mode.replace("_", " ").title()
        items = ch.items or [None]
        for it in items:
            pack = None
            if it and it.packaging_size_id:

                pack = (
                    db.query(PackagingSize).filter(PackagingSize.id == it.packaging_size_id).first()
                )
            rows.append(
                {
                    "Challan Number": ch.challan_number,
                    "Challan Date": ch.challan_date,
                    "Customer Name": ch.customer.name if ch.customer else None,
                    "Product Name": it.product.name if it and it.product else None,
                    "Packaging Size": pack.size_name if pack else None,
                    "Quantity": it.quantity if it else None,
                    "Rate per Kg": it.rate_per_kg if it else None,
                    "Amount": it.amount if it else None,
                    "Paid Amount": ch.paid_amount,
                    "Payment Mode": mode,
                }
            )
    return "Sales Challans", headers, rows


def export_customer_payments(db: Session):
    # Columns mirror the import template one-to-one (see
    # generate_sample_templates.make_customer_payments), so an exported file
    # re-imports as-is. Payment Mode is titled for reading; the importer
    # parses modes fuzzily, so "Bank Transfer" round-trips to bank_transfer.
    headers = [
        "Challan Number",
        "Payment Date",
        "Amount",
        "Payment Mode",
        "Transaction Ref",
        "Payment Notes",
    ]
    rows = [
        {
            "Challan Number": p.challan.challan_number if p.challan else None,
            "Payment Date": p.payment_date,
            "Amount": p.amount,
            "Payment Mode": _title(p.payment_mode),
            "Transaction Ref": p.transaction_ref,
            "Payment Notes": p.notes,
        }
        for p in db.query(CustomerPayment).order_by(CustomerPayment.payment_date).all()
    ]
    return "Customer Payments", headers, rows


def export_expenses(db: Session):
    headers = ["Date", "Category", "Description", "Amount", "Payment Mode", "Payee", "Notes"]
    rows = [
        {
            "Date": e.date,
            "Category": e.category.name if e.category else None,
            "Description": e.description,
            "Amount": e.amount,
            "Payment Mode": _title(e.payment_mode),
            "Payee": e.payee,
            "Notes": e.notes,
        }
        for e in db.query(Expense).order_by(Expense.date).all()
    ]
    return "Expenses", headers, rows


def export_employees(db: Session):
    headers = [
        "Employee ID",
        "Employee Name",
        "Employee Type",
        "Phone",
        "Email",
        "Address",
        "Joining Date",
        "Salary",
        "Bank Account No",
        "IFSC",
        "PAN",
        "Aadhaar",
        "Status",
        "Notes",
    ]
    rows = [
        {
            "Employee ID": e.employee_id,
            "Employee Name": e.name,
            "Employee Type": _title(e.employee_type),
            "Phone": e.phone,
            "Email": e.email,
            "Address": e.address,
            "Joining Date": e.joining_date,
            "Salary": e.salary,
            "Bank Account No": e.bank_account_no,
            "IFSC": e.ifsc_code,
            "PAN": e.pan_no,
            "Aadhaar": e.aadhaar_no,
            "Status": _title(e.status),
            "Notes": e.notes,
        }
        for e in db.query(Employee).order_by(Employee.name).all()
    ]
    return "Employees", headers, rows


EXPORTERS = {
    "customers": export_customers,
    "suppliers": export_suppliers,
    "raw_materials": export_raw_materials,
    "products": export_products,
    "purchase_invoices": export_purchase_invoices,
    "sales_challans": export_sales_challans,
    "customer_payments": export_customer_payments,
    "expenses": export_expenses,
    "employees": export_employees,
}


def build_export(db: Session, export_type: str) -> dict[str, Any]:
    """The raw (title, headers, rows) for an export type.

    This is the single source both renderers read from. The .xlsx and the .pdf
    are produced from THIS, never from separate queries, so the two formats
    cannot drift apart - a PDF that disagrees with its own spreadsheet is worse
    than either being wrong alone.
    """
    if export_type not in EXPORTERS:
        raise ValueError(
            f"Unknown export type: {export_type}. Supported: {', '.join(sorted(EXPORTERS.keys()))}"
        )
    title, headers, rows = EXPORTERS[export_type](db)
    return {"export_type": export_type, "title": title, "headers": headers, "rows": rows}


def run_export(db: Session, export_type: str) -> dict[str, Any]:
    """Build the .xlsx bytes for an export type (same keys as IMPORTERS)."""
    data = build_export(db, export_type)
    return {
        "filename": f"{export_type}_export.xlsx",
        "content": _workbook(data["title"], data["headers"], data["rows"]),
        "rows": len(data["rows"]),
        "headers": data["headers"],
    }
