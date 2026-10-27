"""Generate sample Excel template files for each import type.

Run: python scripts/generate_sample_templates.py
Output: templates/ directory with one .xlsx per import type + a README.txt
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
WARN_FONT = Font(italic=True, color="C00000")


def write_sheet(ws, data_rows):
    """Write data_rows (list of dicts) to the worksheet with a styled header."""
    if not data_rows:
        return
    headers = list(data_rows[0].keys())
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r, row in enumerate(data_rows, 2):
        for col, header in enumerate(headers, 1):
            ws.cell(row=r, column=col, value=row.get(header))
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18


def add_notes_sheet(wb, notes):
    """Put usage instructions on a separate sheet so they don't get imported as data."""
    ws = wb.create_sheet("Instructions")
    ws["A1"] = "How to use this template"
    ws["A1"].font = Font(bold=True, size=13)
    for i, note in enumerate(notes, 2):
        cell = ws.cell(row=i, column=1, value=note)
        cell.font = WARN_FONT
    ws.column_dimensions["A"].width = 110


def make_customers():
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    rows = [
        {
            "Customer ID": "CUST-1001",
            "Customer Name": "Anil Gupta",
            "Company Name": "Gupta Provision Store",
            "GST No": "27ABCDE1234F1Z5",
            "PAN No": "ABCDE1234F",
            "Contact Number": "9876543210",
            "Contact Number 2": "",
            "Email": "anil@test.com",
            "Address": "Nashik",
            "Bank Account": "12345678901",
            "IFSC": "SBIN0001234",
            "Bank Name": "SBI",
            "Opening Balance": 5000,
            "Outstanding Balance": 5000,
            "Notes": "",
        },
        {
            "Customer ID": "CUST-1002",
            "Customer Name": "Priya Nair",
            "Company Name": "",
            "GST No": "",
            "PAN No": "",
            "Contact Number": "9876543211",
            "Email": "",
            "Address": "Kochi",
            "Opening Balance": 2500.50,
            "Outstanding Balance": 2500.50,
            "Notes": "",
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "Mandatory columns: Customer Name, Contact Number.",
            "Customer ID is optional; auto-generated if blank.",
            "Existing customer (matched by ID/name/contact) is skipped "
            "unless Opening Balance is given.",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


def make_suppliers():
    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    rows = [
        {
            "Supplier ID": "SUP-2001",
            "Supplier Name": "Spice Wholesale",
            "Company Name": "Spice Wholesale Co.",
            "GST No": "27XYZ9876A1Z2",
            "PAN No": "",
            "Contact Number": "9123456780",
            "Contact Number 2": "",
            "Email": "",
            "Address": "Indore",
            "Bank Account": "9988776655",
            "IFSC": "SBIN0004567",
            "Bank Name": "SBI",
            "Opening Balance": 0,
            "Outstanding Balance": 0,
            "Notes": "",
        },
        {
            "Supplier ID": "SUP-2002",
            "Supplier Name": "Malwa Agro Foods",
            "Company Name": "Malwa Agro Foods",
            "GST No": "",
            "PAN No": "",
            "Contact Number": "9123456781",
            "Email": "",
            "Address": "Ujjain",
            "Opening Balance": 0,
            "Outstanding Balance": 0,
            "Notes": "",
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "Mandatory columns: Supplier Name, Contact Number.",
            "Supplier ID is optional; auto-generated if blank.",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


def make_raw_materials():
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Materials"
    rows = [
        {
            "Raw Material ID": "RM-3001",
            "Raw Material Name": "Guntur Chilli",
            "Category": "Spices",
            "Unit": "KG",
            "Opening Stock": 100,
            "Threshold Stock": 20,
            "Purchase Rate": 120,
            "HSN Code": "09042110",
        },
        {
            "Raw Material ID": "RM-3002",
            "Raw Material Name": "Turmeric",
            "Category": "Spices",
            "Unit": "KG",
            "Opening Stock": 80,
            "Threshold Stock": 15,
            "Purchase Rate": 90,
            "HSN Code": "09103020",
        },
        {
            "Raw Material ID": "RM-3003",
            "Raw Material Name": "Coriander",
            "Category": "Spices",
            "Unit": "KG",
            "Opening Stock": 60,
            "Threshold Stock": 10,
            "Purchase Rate": 70,
            "HSN Code": "09092100",
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "Mandatory column: Raw Material Name.",
            "Unit accepts: kg, gram, packet, pieces, litres, metres "
            "(plus common aliases like pcs, ltr, mtr). Category is auto-created if new.",
            "HSN Code is optional, digits only; pick it from the HSN dropdown in Inventory.",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


def make_products():
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    rows = [
        {
            "Product ID": "PROD-4001",
            "Product Name": "Shresth Chilli Masala",
            "Category": "Dry Masala",
            "Description": "",
            "Rate per Kg": 240,
            "Packaging Size": "50g",
            "Weight (kg)": 0.05,
            "HSN Code": "09109990",
        },
        {
            "Product ID": "PROD-4001",
            "Product Name": "Shresth Chilli Masala",
            "Category": "Dry Masala",
            "Description": "",
            "Rate per Kg": 240,
            "Packaging Size": "100g",
            "Weight (kg)": 0.1,
            "HSN Code": "09109990",
        },
        {
            "Product ID": "PROD-4002",
            "Product Name": "Shresth Turmeric Powder",
            "Category": "Dry Masala",
            "Rate per Kg": 140,
            "Packaging Size": "100g",
            "Weight (kg)": 0.1,
            "HSN Code": "09103020",
        },
        {
            "Product ID": "PROD-4002",
            "Product Name": "Shresth Turmeric Powder",
            "Category": "Dry Masala",
            "Rate per Kg": 140,
            "Packaging Size": "250g",
            "Weight (kg)": 0.25,
            "HSN Code": "09103020",
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "Mandatory column: Product Name.",
            "Package several rows with the same Product Name to add multiple packing sizes.",
            "Packaging Size can be text like '50g' - weight is parsed "
            "from it if Weight (kg) is blank.",
            "HSN Code is optional, digits only; pick it from the HSN dropdown in Inventory.",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


def make_purchase_invoices():
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Invoices"
    rows = [
        {
            "Invoice Number": "PUR-2026-001",
            "Invoice Date": date(2026, 9, 8),
            "Supplier Name": "Spice Wholesale",
            "Raw Material Name": "Guntur Chilli",
            "Unit": "KG",
            "Quantity": 50,
            "Rate": 110,
            "GST %": 5,
            "Paid Amount": 0,
        },
        {
            "Invoice Number": "PUR-2026-001",
            "Invoice Date": date(2026, 9, 8),
            "Supplier Name": "Spice Wholesale",
            "Raw Material Name": "Turmeric",
            "Unit": "KG",
            "Quantity": 30,
            "Rate": 85,
            "GST %": 5,
            "Paid Amount": 0,
        },
        {
            "Invoice Number": "PUR-2026-002",
            "Invoice Date": date(2026, 9, 12),
            "Supplier Name": "Malwa Agro Foods",
            "Raw Material Name": "Coriander",
            "Unit": "KG",
            "Quantity": 40,
            "Rate": 68,
            "GST %": 5,
            "Paid Amount": 0,
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "Repeat the same Invoice Number across multiple rows for multiple materials.",
            "Mandatory: Supplier Name, Raw Material Name (supplier matches name "
            "or company name - import suppliers/raw_materials first).",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


def make_sales_challans():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Challans"
    rows = [
        {
            "Challan Number": "CH-2026-001",
            "Challan Date": date(2026, 9, 20),
            "Customer Name": "Anil Gupta",
            "Product Name": "Shresth Chilli Masala",
            "Packaging Size": "50g",
            "Quantity": 100,
            "Rate per Kg": 240,
            "Amount": 1200,
            "Paid Amount": 1200,
            "Payment Mode": "UPI",
        },
        {
            "Challan Number": "CH-2026-002",
            "Challan Date": date(2026, 9, 22),
            "Customer Name": "Priya Nair",
            "Product Name": "Shresth Turmeric Powder",
            "Packaging Size": "100g",
            "Quantity": 50,
            "Rate per Kg": 140,
            "Amount": 700,
            "Paid Amount": 0,
        },
        {
            "Challan Number": "CH-2026-002",
            "Challan Date": date(2026, 9, 22),
            "Customer Name": "Priya Nair",
            "Product Name": "Shresth Chilli Masala",
            "Packaging Size": "100g",
            "Quantity": 20,
            "Rate per Kg": 240,
            "Amount": 480,
            "Paid Amount": 0,
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "Repeat the same Challan Number to add more line items to one challan.",
            "Mandatory: Customer Name (matched by name - import customers first; "
            "products are auto-created if new).",
            "If Amount is blank it is computed as Weight x Rate per Kg.",
            "Customer/Product matched by name. If the challan is only partly paid, "
            "outstanding is tracked.",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


def make_customer_payments():
    wb = Workbook()
    ws = wb.active
    ws.title = "Customer Payments"
    rows = [
        {
            "Challan Number": "CH-2026-002",
            "Payment Date": date(2026, 9, 28),
            "Amount": 500,
            "Payment Mode": "Cash",
            "Transaction Ref": "",
            "Payment Notes": "Part payment",
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "One row = one payment. Repeat the Challan Number for part-payments.",
            "Mandatory: Challan Number (must already exist - import sales "
            "challans first), Amount.",
            "Draft or archived challans are skipped, as is any amount above "
            "the challan's outstanding.",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


def make_expenses():
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    rows = [
        {
            "Date": date(2026, 9, 15),
            "Category": "Electricity",
            "Description": "September electricity bill",
            "Amount": 1500,
            "Payment Mode": "UPI",
            "Payee": "MSEB",
            "Notes": "",
        },
        {
            "Date": date(2026, 9, 25),
            "Category": "Fuel",
            "Description": "Diesel for delivery van",
            "Amount": 2200,
            "Payment Mode": "Cash",
            "Payee": "Fuel station",
            "Notes": "",
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "Mandatory: Description and Amount (a row with both blank is skipped).",
            "Category is auto-created if new. Missing category defaults to Miscellaneous.",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


def make_employees():
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    rows = [
        {
            "Employee ID": "EMP-1001",
            "Employee Name": "Rakesh Verma",
            "Employee Type": "Full Time",
            "Phone": "9823456710",
            "Email": "rakesh@test.com",
            "Address": "Indore",
            "Joining Date": date(2022, 6, 1),
            "Salary": 15000,
            "Bank Account No": "2233445566",
            "IFSC": "SBIN0009999",
            "PAN": "AAAAA1111A",
            "Aadhaar": "111122223333",
            "Status": "Active",
            "Notes": "",
        },
        {
            "Employee ID": "EMP-1002",
            "Employee Name": "Sunita Yadav",
            "Employee Type": "Part Time",
            "Phone": "9823456711",
            "Email": "",
            "Address": "Indore",
            "Joining Date": date(2023, 4, 1),
            "Salary": 8000,
            "Bank Account No": "",
            "IFSC": "",
            "PAN": "",
            "Aadhaar": "444455556666",
            "Status": "Active",
            "Notes": "",
        },
    ]
    write_sheet(ws, rows)
    add_notes_sheet(
        wb,
        [
            "Mandatory column: Employee Name.",
            "Employee Type accepts: Full Time / Part Time / Contractor.",
            "Delete the sample rows below the header before importing your own data.",
        ],
    )
    return wb


FACTORIES = {
    "customers": make_customers,
    "suppliers": make_suppliers,
    "raw_materials": make_raw_materials,
    "products": make_products,
    "purchase_invoices": make_purchase_invoices,
    "sales_challans": make_sales_challans,
    "customer_payments": make_customer_payments,
    "expenses": make_expenses,
    "employees": make_employees,
}


def main():
    out_dir = Path(__file__).resolve().parent.parent / "templates"
    out_dir.mkdir(exist_ok=True)

    for name, factory in FACTORIES.items():
        wb = factory()
        path = out_dir / f"sample_{name}.xlsx"
        wb.save(path)
        print(f"Wrote {path}")

    readme = out_dir / "README.txt"
    readme.write_text(
        "SAMPLE IMPORT TEMPLATES\n"
        "=======================\n"
        "\n"
        "Each file's first row holds the column headers - do not rename or remove them.\n"
        "Delete the sample data rows and add your own; the header row must stay in row 1.\n"
        "\n"
        "Import order (important):\n"
        "  1. customers\n"
        "  2. suppliers\n"
        "  3. raw_materials\n"
        "  4. products\n"
        "  5. purchase_invoices (needs suppliers + raw_materials)\n"
        "  6. sales_challans (needs customers + products)\n"
        "  7. customer_payments (needs sales challans)\n"
        "  8. expenses\n"
        "  9. employees\n"
        "\n"
        "General tips:\n"
        "- Dates work in many formats: 2024-01-05, 05-01-2024, 05/01/2024\n"
        "- Amounts tolerate currency symbols and commas (e.g. Rs 1,500)\n"
        "- Coded IDs are optional - the system generates them when left blank.\n"
        "- Re-running an import skips rows that already exist "
        "(matches on code/name/challan number).\n",
        encoding="utf-8",
    )
    print(f"Wrote {readme}")


if __name__ == "__main__":
    main()
